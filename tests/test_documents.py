"""Tests des outils Word, sur des documents fabriqués dans un dossier temporaire.

Ce qui compte le plus : l'original n'est jamais modifié, un nom venu du
modèle ne sort jamais du dossier, et l'écriture ne passe pas sans validation.
"""

from pathlib import Path

import docx
import pytest

from agent.approval import DenyAll
from agent.documents import (
    DocumentFolder,
    buildDocumentTools,
    createDocumentsFromSpec,
    replaceText,
)
from agent.errors import DocumentFolderError, ToolError
from agent.loop import runAgent
from agent.tools import ToolRegistry
from tests.conftest import ScriptedClient, textTurn, toolTurn

DEMO_DOCUMENTS = Path("evaluations/jeux/documents_demo.jsonl")


@pytest.fixture
def folder(tmp_path: Path) -> DocumentFolder:
    """Un dossier d'entrée avec un document fictif, et un dossier de sortie."""
    inputDir = tmp_path / "entree"
    inputDir.mkdir()
    document = docx.Document()
    document.add_heading("Compte rendu", level=1)
    paragraph = document.add_paragraph()
    # « jeudi » coupé sur deux runs, le premier en gras : c'est ce que Word
    # produit après une correction orthographique.
    paragraph.add_run("Réunion le jeu").bold = True
    paragraph.add_run("di à 10 heures.")
    document.add_paragraph("Fin du document.")
    document.save(str(inputDir / "cr.docx"))
    (inputDir / "notes.txt").write_text("pas un document", encoding="utf-8")
    return DocumentFolder(inputDir, tmp_path / "sortie")


@pytest.fixture
def tools(folder: DocumentFolder) -> ToolRegistry:
    """Registre des trois outils Word."""
    return ToolRegistry(buildDocumentTools(folder))


def call(registry: ToolRegistry, name: str, **arguments: object) -> str:
    """Appelle directement le gestionnaire d'un outil."""
    return registry.get(name).handler(arguments)


def test_onlyWordDocumentsAreListed(tools: ToolRegistry) -> None:
    assert call(tools, "listDocuments") == "cr.docx"


def test_readingGivesNumberedParagraphsWithStyle(tools: ToolRegistry) -> None:
    text = call(tools, "readDocument", nom="cr.docx")

    assert "[0] (Heading 1) Compte rendu" in text
    assert "[1] (Normal) Réunion le jeudi à 10 heures." in text


def test_unknownDocumentListsTheAvailableOnes(tools: ToolRegistry) -> None:
    # Le modèle doit pouvoir se rattraper en un tour au lieu de réessayer.
    with pytest.raises(ToolError, match=r"cr\.docx"):
        call(tools, "readDocument", nom="inconnu.docx")


def test_replacementAcrossRunsKeepsFormatting(
    tools: ToolRegistry, folder: DocumentFolder
) -> None:
    call(
        tools,
        "writeDocumentCopy",
        nom="cr.docx",
        nomCopie="cr-modifie.docx",
        remplacements=[{"ancien": "jeudi", "nouveau": "vendredi"}],
    )

    copy = docx.Document(str(folder.outputDir / "cr-modifie.docx"))
    paragraph = copy.paragraphs[1]
    assert paragraph.text == "Réunion le vendredi à 10 heures."
    # Le texte remplacé prend la mise en forme du premier run concerné.
    assert paragraph.runs[0].text == "Réunion le vendredi"
    assert paragraph.runs[0].bold is True
    assert paragraph.runs[1].bold is None


def test_originalIsNeverModified(tools: ToolRegistry, folder: DocumentFolder) -> None:
    call(
        tools,
        "writeDocumentCopy",
        nom="cr.docx",
        nomCopie="cr-modifie.docx",
        remplacements=[{"ancien": "jeudi", "nouveau": "vendredi"}],
    )

    original = docx.Document(str(folder.inputDir / "cr.docx"))
    assert "jeudi" in original.paragraphs[1].text


def test_existingCopyIsNeverOverwritten(tools: ToolRegistry) -> None:
    arguments = {
        "nom": "cr.docx",
        "nomCopie": "cr-modifie.docx",
        "remplacements": [{"ancien": "jeudi", "nouveau": "vendredi"}],
    }
    call(tools, "writeDocumentCopy", **arguments)

    with pytest.raises(ToolError, match="existe déjà"):
        call(tools, "writeDocumentCopy", **arguments)


def test_missingTextWritesNoCopy(tools: ToolRegistry, folder: DocumentFolder) -> None:
    with pytest.raises(ToolError, match="introuvable"):
        call(
            tools,
            "writeDocumentCopy",
            nom="cr.docx",
            nomCopie="cr-modifie.docx",
            remplacements=[{"ancien": "lundi", "nouveau": "mardi"}],
        )

    assert not (folder.outputDir / "cr-modifie.docx").exists()


def test_namesCannotLeaveTheFolder(tools: ToolRegistry) -> None:
    # Des noms hostiles, pas seulement malformés : un nom vient du modèle,
    # donc potentiellement d'un document qu'il a lu.
    for name in ["../cr.docx", "sous/cr.docx", "/tmp/cr.docx", ".cache.docx", "cr"]:
        with pytest.raises(ToolError, match="invalide"):
            call(tools, "readDocument", nom=name)
        with pytest.raises(ToolError, match="invalide"):
            call(
                tools,
                "writeDocumentCopy",
                nom="cr.docx",
                nomCopie=name,
                remplacements=[{"ancien": "jeudi", "nouveau": "vendredi"}],
            )


def test_sectionIsInsertedAfterTheGivenParagraph(
    tools: ToolRegistry, folder: DocumentFolder
) -> None:
    call(
        tools,
        "writeDocumentCopy",
        nom="cr.docx",
        nomCopie="cr-section.docx",
        sections=[
            {"titre": "Décisions", "paragraphes": ["Une", "Deux"], "apres": 1},
            {"titre": "Annexe", "paragraphes": ["Fin."], "niveau": 2},
        ],
    )

    copy = docx.Document(str(folder.outputDir / "cr-section.docx"))
    summary = [(paragraph.style.name, paragraph.text) for paragraph in copy.paragraphs]
    assert summary[2:5] == [
        ("Heading 1", "Décisions"),
        ("Normal", "Une"),
        ("Normal", "Deux"),
    ]
    assert summary[5] == ("Normal", "Fin du document.")
    assert summary[-2:] == [("Heading 2", "Annexe"), ("Normal", "Fin.")]


def test_invalidSectionIsRefused(tools: ToolRegistry) -> None:
    for section in [
        {"titre": "X", "paragraphes": [], "apres": 99},
        {"titre": "X", "paragraphes": [], "niveau": 7},
        {"titre": "", "paragraphes": ["texte"]},
        {"titre": "X", "paragraphes": "pas une liste"},
    ]:
        with pytest.raises(ToolError):
            call(
                tools,
                "writeDocumentCopy",
                nom="cr.docx",
                nomCopie="cr-x.docx",
                sections=[section],
            )


def test_emptyRequestIsRefused(tools: ToolRegistry) -> None:
    with pytest.raises(ToolError, match="Aucune modification"):
        call(tools, "writeDocumentCopy", nom="cr.docx", nomCopie="cr-x.docx")


def test_replacementContainingTheOldTextTerminates() -> None:
    paragraph = docx.Document().add_paragraph("jeudi et jeudi")

    assert replaceText(paragraph, "jeudi", "jeudi 12") == 2
    assert paragraph.text == "jeudi 12 et jeudi 12"


def test_sameInputAndOutputFolderIsRefused(tmp_path: Path) -> None:
    with pytest.raises(DocumentFolderError, match="différent"):
        DocumentFolder(tmp_path, tmp_path)


def test_writingIsSensitiveAndSaysSo(tools: ToolRegistry) -> None:
    assert tools.sensitiveNames() == ["writeDocumentCopy"]
    description = tools.get("writeDocumentCopy").description.lower()
    assert "irréversible" in description
    assert "validation" in description


def test_refusedWritingWritesNothing(
    tools: ToolRegistry, folder: DocumentFolder
) -> None:
    client = ScriptedClient(
        [
            toolTurn(
                (
                    "c1",
                    "writeDocumentCopy",
                    {
                        "nom": "cr.docx",
                        "nomCopie": "cr-modifie.docx",
                        "remplacements": [{"ancien": "jeudi", "nouveau": "vendredi"}],
                    },
                )
            ),
            textTurn("L'écriture a été refusée faute de validation."),
        ]
    )

    trace = runAgent(
        client=client,
        registry=tools,
        approval=DenyAll(),
        systemPrompt="Consignes de test.",
        task="modifie le document",
    )

    assert trace.refusedTools() == ["writeDocumentCopy"]
    assert not folder.outputDir.exists()


def test_demoDocumentsAreBuilt(tmp_path: Path) -> None:
    createDocumentsFromSpec(DEMO_DOCUMENTS, tmp_path)

    folder = DocumentFolder(tmp_path, tmp_path / "sortie")
    assert folder.listNames() == [
        "compte-rendu-reunion.docx",
        "consignes-diffusion.docx",
    ]
