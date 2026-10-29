"""Tests de l'interface web.

Sans LLM : l'agent est joué par `ScriptedClient`. Le serveur HTTP est lancé
pour de vrai, sur un port libre de 127.0.0.1, pour vérifier les adresses.
"""

import threading
import time
import urllib.error
import urllib.request
from http.server import ThreadingHTTPServer
from pathlib import Path

import docx
import pytest

from agent.errors import UploadError
from agent.models import ToolCall
from agent.web import (
    STATUS_DONE,
    Handler,
    WebApp,
    WebTask,
    describeCall,
    parseUpload,
    runTask,
    safeDocumentName,
)
from tests.conftest import ScriptedClient, textTurn, toolTurn

BOUNDARY = "frontiere"


def multipart(fields: dict[str, tuple[str | None, bytes]]) -> bytes:
    """Corps de formulaire tel qu'un navigateur l'envoie.

    Chaque champ : (nom de fichier ou None, contenu).
    """
    body = b""
    for name, (filename, content) in fields.items():
        disposition = f'form-data; name="{name}"'
        if filename:
            disposition += f'; filename="{filename}"'
        body += (
            f"--{BOUNDARY}\r\nContent-Disposition: {disposition}\r\n\r\n".encode()
            + content
            + b"\r\n"
        )
    return body + f"--{BOUNDARY}--\r\n".encode()


CONTENT_TYPE = f"multipart/form-data; boundary={BOUNDARY}"


def wordBytes(tmp_path: Path) -> bytes:
    """Un petit document Word fictif, en octets."""
    document = docx.Document()
    document.add_paragraph("Réunion jeudi à 10 heures.")
    path = tmp_path / "source.docx"
    document.save(str(path))
    return path.read_bytes()


def test_uploadIsParsed() -> None:
    body = multipart(
        {
            "document": ("compte-rendu.docx", b"octets"),
            "demande": (None, "remplace « jeudi »".encode()),
        }
    )
    assert parseUpload(CONTENT_TYPE, body) == (
        "compte-rendu.docx",
        b"octets",
        "remplace « jeudi »",
    )


def test_uploadWithoutInstructionIsRefused() -> None:
    body = multipart({"document": ("a.docx", b"x"), "demande": (None, b"  ")})
    with pytest.raises(UploadError):
        parseUpload(CONTENT_TYPE, body)


def test_uploadedNameLosesItsFolder() -> None:
    assert safeDocumentName("C:\\Users\\moi\\Rapport.DOCX") == "Rapport.docx"
    assert safeDocumentName("../../etc/a.docx") == "a.docx"
    for name in ["a.pdf", ".cache.docx", "~$brouillon.docx", "docx"]:
        with pytest.raises(UploadError):
            safeDocumentName(name)


def test_proposalIsShownEscaped() -> None:
    # Les arguments viennent du modèle, qui a lu le document : ils peuvent
    # contenir du HTML. Il doit s'afficher comme du texte.
    call = ToolCall(
        callId="c1",
        toolName="writeDocumentCopy",
        arguments={
            "nomCopie": "a.docx",
            "remplacements": [{"ancien": "<b>x</b>", "nouveau": "y"}],
        },
    )
    shown = describeCall(call)
    assert "&lt;b&gt;x&lt;/b&gt;" in shown
    assert "<b>x</b>" not in shown


def runWithDecision(task: WebTask, client: ScriptedClient, approved: bool) -> None:
    """Lance l'agent, puis clique Valider ou Refuser quand il le demande."""
    worker = threading.Thread(
        target=runTask, args=(task, client, "prompt", "remplace jeudi")
    )
    worker.start()
    while task.approval.pendingCall is None:
        time.sleep(0.01)
    task.approval.decide(approved)
    worker.join(timeout=5)


WRITE_TURN = toolTurn(
    (
        "c1",
        "writeDocumentCopy",
        {
            "nom": "cr.docx",
            "nomCopie": "cr-modifie.docx",
            "remplacements": [{"ancien": "jeudi", "nouveau": "vendredi"}],
        },
    )
)


def newTask(tmp_path: Path) -> WebTask:
    task = WebTask("t1", tmp_path / "t1", "cr.docx")
    (task.folder.inputDir / "cr.docx").write_bytes(wordBytes(tmp_path))
    return task


def test_validatedTaskGivesACopy(tmp_path: Path) -> None:
    task = newTask(tmp_path)
    original = (task.folder.inputDir / "cr.docx").read_bytes()

    runWithDecision(task, ScriptedClient([WRITE_TURN, textTurn("Fait.")]), True)

    assert task.status == STATUS_DONE
    assert task.copies() == ["cr-modifie.docx"]
    copy = docx.Document(str(task.folder.outputDir / "cr-modifie.docx"))
    assert copy.paragraphs[0].text == "Réunion vendredi à 10 heures."
    assert (task.folder.inputDir / "cr.docx").read_bytes() == original


def test_refusedTaskWritesNothing(tmp_path: Path) -> None:
    task = newTask(tmp_path)

    runWithDecision(task, ScriptedClient([WRITE_TURN, textTurn("Refusé.")]), False)

    assert task.status == STATUS_DONE
    assert task.copies() == []


@pytest.fixture
def server(tmp_path: Path):
    """Serveur réel sur un port libre ; l'agent n'est jamais lancé ici."""
    Handler.app = WebApp(ScriptedClient([]), "prompt", tmp_path / "interface")
    httpd = ThreadingHTTPServer(("127.0.0.1", 0), Handler)
    thread = threading.Thread(target=httpd.serve_forever, daemon=True)
    thread.start()
    yield f"http://127.0.0.1:{httpd.server_address[1]}", Handler.app
    httpd.shutdown()
    httpd.server_close()


def test_formIsServed(server) -> None:
    url, _ = server
    with urllib.request.urlopen(url + "/") as response:
        assert 'type="file"' in response.read().decode("utf-8")


def test_downloadServesOnlyWrittenCopies(server, tmp_path: Path) -> None:
    url, app = server
    task = newTask(tmp_path)
    app.tasks[task.taskId] = task
    task.folder.outputDir.mkdir()
    (task.folder.outputDir / "cr-modifie.docx").write_bytes(b"copie")

    with urllib.request.urlopen(f"{url}/taches/t1/copie/cr-modifie.docx") as response:
        assert response.read() == b"copie"
    # L'original, ou tout chemin fabriqué, n'est pas servi.
    for path in ["cr.docx", "..%2Fentree%2Fcr.docx", "inconnu.docx"]:
        with pytest.raises(urllib.error.HTTPError) as error:
            urllib.request.urlopen(f"{url}/taches/t1/copie/{path}")
        assert error.value.code == 404


def test_unknownTaskIs404(server) -> None:
    url, _ = server
    with pytest.raises(urllib.error.HTTPError) as error:
        urllib.request.urlopen(url + "/taches/inconnue")
    assert error.value.code == 404
